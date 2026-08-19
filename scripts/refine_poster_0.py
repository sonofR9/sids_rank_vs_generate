from __future__ import annotations

# Hallmark · pre-emit critique: P5 H5 E4 S5 R5 V4
# Constrained Narrative Workflow · technical/austere · inherited poster_0 system

from hashlib import sha256
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_poster_pptx import (
    GREEN,
    GREEN_DARK,
    INK,
    MUTED,
    PURPLE_RULE,
    PURPLE_TINT,
    add_band,
    add_flow_arrow,
    add_rich_textbox,
    add_rule,
    add_textbox,
)


SOURCE = Path("/Users/konstantinryadinskiy/Desktop/poster_0.pptx")
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_0_improved.pptx"
)

EMU_PER_INCH = 914400
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def delete_named_content(slide, prefixes: tuple[str, ...], *, keep: set[str] | None = None) -> None:
    keep = keep or set()
    for shape in list(slide.shapes):
        if shape.name in keep:
            continue
        if shape.name.startswith(prefixes):
            shape._element.getparent().remove(shape._element)


def delete_hidden_named_content(slide, prefixes: tuple[str, ...]) -> None:
    """Remove Office Math AlternateContent wrappers from the replaced content layer."""
    c_nv_pr_tag = f"{{{P_NS}}}cNvPr"
    alternate_tag = f"{{{MC_NS}}}AlternateContent"
    wrappers = set()
    for node in slide._element.iter(c_nv_pr_tag):
        if not node.get("name", "").startswith(prefixes):
            continue
        parent = node.getparent()
        while parent is not None and parent.tag != alternate_tag:
            parent = parent.getparent()
        if parent is not None:
            wrappers.add(parent)
    for wrapper in wrappers:
        wrapper.getparent().remove(wrapper)


def add_inline_statement(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    body: str,
    *,
    name: str,
    label_color=MUTED,
    size: float = 30,
):
    return add_rich_textbox(
        slide,
        x,
        y,
        w,
        h,
        [
            {
                "runs": [
                    (f"{label} · ", size, True, label_color),
                    (body, size, False, INK),
                ],
                "line_spacing": 1.02,
            }
        ],
        name=name,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )


def build_introduction(slide) -> None:
    x, w = 1.86, 12.79

    add_textbox(
        slide,
        x,
        9.50,
        w,
        2.15,
        "Какие свойства Semantic ID действительно отражаются на downstream-метриках "
        "рекомендательной системы?",
        size=40,
        bold=True,
        name="INTRO_ResearchQuestion",
        line_spacing=1.0,
    )
    add_rule(
        slide,
        x,
        11.70,
        x + 7.30,
        11.70,
        color=GREEN,
        width=4,
        name="INTRO_QuestionUnderline",
    )

    add_textbox(
        slide,
        x,
        11.92,
        w,
        0.48,
        "ЧТО ТАКОЕ SID?",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_DefinitionHead",
    )
    add_textbox(
        slide,
        x,
        12.40,
        w,
        1.02,
        "Semantic IDs — дискретные коды объектов для генеративных рекомендательных моделей.",
        size=30,
        bold=True,
        name="INTRO_Definition",
        line_spacing=1.02,
    )

    add_band(slide, x, 13.55, w, 2.42, name="INTRO_RQKMeansBand")
    add_textbox(
        slide,
        x + 0.18,
        13.65,
        w - 0.36,
        0.44,
        "RQ-KMeans · residual quantization",
        size=30,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_RQKMeansHead",
    )
    stage_y, stage_h = 14.48, 0.72
    stages = [
        (0.12, 2.34, "item\nembedding"),
        (2.88, 2.40, "nearest code\nc₁"),
        (5.72, 1.92, "residual"),
        (8.08, 2.42, "next code\nc₂"),
        (10.94, 1.61, "…"),
    ]
    for index, (offset, width, label) in enumerate(stages):
        add_textbox(
            slide,
            x + offset,
            stage_y,
            width,
            stage_h,
            label,
            size=30,
            bold=True,
            color=GREEN_DARK if label == "…" else INK,
            align=PP_ALIGN.CENTER,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"INTRO_RQKMeansStage_{index}",
            line_spacing=1.0,
        )
        if index < len(stages) - 1:
            arrow_x = x + offset + width + 0.10
            add_flow_arrow(
                slide,
                arrow_x,
                stage_y + 0.09,
                w=0.32,
                name=f"INTRO_RQKMeansArrow_{index}",
            )
    add_textbox(
        slide,
        x + 0.20,
        15.40,
        w - 0.40,
        0.62,
        "c₁  →  c₂  →  …  →  cₗ   =   Semantic ID",
        size=34,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansSequence",
    )

    add_textbox(
        slide,
        x,
        16.32,
        w,
        0.48,
        "ЧТО ХОТИМ ВЫЯСНИТЬ?",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_GoalHead",
    )
    add_rich_textbox(
        slide,
        x,
        16.84,
        w,
        1.60,
        [
            {"runs": [("• Сравнить способы построения и модификации SID.", 30, True, INK)]},
            {
                "runs": [("• Проверить связь характеристик SID с downstream-качеством.", 30, True, INK)],
                "before": 5,
            },
        ],
        name="INTRO_GoalBody",
    )

    add_textbox(
        slide,
        x,
        18.62,
        w,
        0.48,
        "ГДЕ ПРОВЕРЯЕМ?",
        size=34,
        bold=True,
        name="INTRO_ContextHead",
    )
    add_textbox(
        slide,
        x,
        19.12,
        w,
        0.52,
        "Yambda-50M · time split",
        size=30,
        bold=True,
        color=MUTED,
        name="INTRO_Dataset",
    )
    add_rule(
        slide,
        x + 6.22,
        19.84,
        x + 6.22,
        24.18,
        color=PURPLE_RULE,
        width=1.5,
        name="INTRO_TaskDivider",
    )
    add_textbox(
        slide,
        x,
        19.84,
        5.92,
        0.50,
        "Generation",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_GenerationHead",
    )
    add_rich_textbox(
        slide,
        x,
        20.42,
        5.92,
        3.50,
        [
            {"runs": [("TIGER → SID", 30, True, INK)]},
            {"runs": [("Baseline · SASRec", 30, False, INK)], "before": 6},
            {"runs": [("Recall@100", 30, False, INK)], "before": 6},
        ],
        name="INTRO_GenerationBody",
    )
    add_textbox(
        slide,
        x + 6.52,
        19.84,
        6.27,
        0.50,
        "Ranking",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_RankingHead",
    )
    add_rich_textbox(
        slide,
        x + 6.52,
        20.42,
        6.27,
        3.64,
        [
            {"runs": [("DCNv2 / TransAct-like", 30, True, INK)]},
            {"runs": [("SID — дополнительный признак", 30, False, INK)], "before": 6},
            {
                "runs": [("ROC-AUC · Pairwise Accuracy", 30, False, INK)],
                "before": 6,
                "line_spacing": 1.02,
            },
        ],
        name="INTRO_RankingBody",
    )


def build_methods(slide, rqvae_picture) -> None:
    x, w = 15.78, 13.78
    half_w, gap = 6.58, 0.62
    right_x = x + half_w + gap

    add_band(slide, x, 9.50, w, 1.30, name="METHODS_BaselineBand")
    add_inline_statement(
        slide,
        x + 0.16,
        9.58,
        half_w - 0.16,
        1.12,
        "Balanced RQ-KMeans",
        "capacity ограничивает размер кластера.",
        name="METHODS_Balanced",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x + half_w + gap / 2,
        9.66,
        x + half_w + gap / 2,
        10.64,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_BaselineDivider",
    )
    add_inline_statement(
        slide,
        right_x,
        9.58,
        half_w,
        1.12,
        "PLUM-like",
        "меньше кодов + случайный активный prefix.",
        name="METHODS_PLUM",
        label_color=GREEN_DARK,
    )

    add_textbox(
        slide,
        x,
        11.02,
        2.20,
        0.50,
        "RQ-VAE",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_RQVAEHead",
    )
    add_textbox(
        slide,
        x + 2.20,
        11.02,
        w - 2.20,
        0.56,
        "Residual quantization с обучаемым encoder / decoder.",
        size=30,
        name="METHODS_RQVAENote",
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )
    rqvae_picture.left = Inches(17.87)
    rqvae_picture.top = Inches(11.52)
    rqvae_picture.width = Inches(9.60)
    rqvae_picture.height = Inches(4.30)

    add_textbox(
        slide,
        x,
        15.96,
        w,
        0.56,
        "User-aware",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_UserAwareHead",
    )
    add_inline_statement(
        slide,
        x,
        16.56,
        w,
        0.70,
        "Что меняем",
        "учитываем пользовательские предпочтения при построении SID.",
        name="METHODS_UserAwareWhat",
    )
    add_band(slide, x, 17.34, w, 1.22, name="METHODS_UserAwareBand")
    user_stages = [
        (0.05, 3.05, "USER + ITEM\nEMBEDDING"),
        (3.55, 3.40, "USER-AWARE\nMODEL"),
        (7.40, 2.70, "U-AWARE\nSID"),
        (10.55, 3.05, "TARGET SID"),
    ]
    for index, (offset, width, label) in enumerate(user_stages):
        add_textbox(
            slide,
            x + offset,
            17.48,
            width,
            0.92,
            label,
            size=30,
            bold=True,
            color=GREEN_DARK if index == 2 else INK,
            align=PP_ALIGN.CENTER,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"METHODS_UserAwareStage_{index}",
            line_spacing=1.0,
        )
        if index < len(user_stages) - 1:
            add_flow_arrow(
                slide,
                x + offset + width + 0.08,
                17.64,
                w=0.34,
                name=f"METHODS_UserAwareArrow_{index}",
            )
    add_inline_statement(
        slide,
        x,
        18.64,
        w,
        0.62,
        "Зачем",
        "отразить пользовательское восприятие, не только content similarity.",
        name="METHODS_UserAwareWhy",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        19.30,
        x + w,
        19.30,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_UserAwareRule",
    )

    add_textbox(
        slide,
        x,
        19.46,
        w,
        0.56,
        "PSI",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_PSIHead",
    )
    add_inline_statement(
        slide,
        x,
        20.06,
        w,
        0.70,
        "Что меняем",
        "занятый nearest SID → следующая близкая свободная комбинация.",
        name="METHODS_PSIWhat",
    )
    add_band(slide, x, 20.82, w, 1.10, name="METHODS_PSIBand")
    psi_stages = [
        (0.16, 3.20, "nearest SID\nзанят  ×", MUTED),
        (3.84, 3.28, "next closest\nзанят  ×", MUTED),
        (7.60, 6.02, "next closest free SID  ✓", GREEN_DARK),
    ]
    for index, (offset, width, label, color) in enumerate(psi_stages):
        add_textbox(
            slide,
            x + offset,
            20.96,
            width,
            0.82,
            label,
            size=30,
            bold=True,
            color=color,
            align=PP_ALIGN.CENTER,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"METHODS_PSIStage_{index}",
            line_spacing=1.0,
        )
        if index < len(psi_stages) - 1:
            add_flow_arrow(
                slide,
                x + offset + width + 0.08,
                21.10,
                w=0.34,
                name=f"METHODS_PSIArrow_{index}",
            )
    add_inline_statement(
        slide,
        x,
        22.00,
        w,
        0.62,
        "Зачем",
        "устранить SID-коллизии без дополнительного токена.",
        name="METHODS_PSIWhy",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        22.68,
        x + w,
        22.68,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_PSIRule",
    )

    add_textbox(
        slide,
        x,
        22.84,
        half_w,
        0.50,
        "R3-VAE",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_R3VAEHead",
    )
    add_inline_statement(
        slide,
        x,
        23.46,
        half_w,
        0.92,
        "Что меняем",
        "semantic anchor перед residual Q.",
        name="METHODS_R3VAEWhat",
    )
    add_band(slide, x, 24.42, half_w, 0.90, name="METHODS_R3VAEBand")
    add_textbox(
        slide,
        x + 0.12,
        24.54,
        half_w - 0.24,
        0.64,
        "embedding + anchor  →  residual Q",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_R3VAEScheme",
    )
    add_inline_statement(
        slide,
        x,
        25.42,
        half_w,
        0.60,
        "Зачем",
        "структурировать пространство кодов.",
        name="METHODS_R3VAEWhy",
        label_color=GREEN_DARK,
    )

    add_textbox(
        slide,
        right_x,
        22.84,
        half_w,
        0.50,
        "LogQ",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_LogQHead",
    )
    add_inline_statement(
        slide,
        right_x,
        23.46,
        half_w,
        0.92,
        "Что меняем",
        "учёт частоты SID-префиксов / item ID.",
        name="METHODS_LogQWhat",
    )
    add_band(slide, right_x, 24.42, half_w, 0.90, name="METHODS_LogQBand")
    add_textbox(
        slide,
        right_x + 0.12,
        24.50,
        half_w - 0.24,
        0.72,
        "logits  →  frequency correction\n→  corrected logits",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_LogQScheme",
        line_spacing=1.0,
    )
    add_inline_statement(
        slide,
        right_x,
        25.42,
        half_w,
        0.60,
        "Зачем",
        "компенсировать frequency bias.",
        name="METHODS_LogQWhy",
        label_color=GREEN_DARK,
    )


def build_generation_table(slide) -> None:
    x, w = 31.19, 13.78
    col_model, col_recall = 7.25, 3.35
    col_delta = w - col_model - col_recall
    header_y, header_h, row_h = 19.14, 0.70, 0.94

    add_textbox(
        slide,
        x,
        18.48,
        w,
        0.50,
        "GENERATION · ОСНОВНОЙ РЕЗУЛЬТАТ",
        size=34,
        bold=True,
        name="RESULTS_GenerationTableHead",
    )
    add_band(slide, x, header_y, w, header_h, fill=PURPLE_TINT, name="RESULTS_TableHeaderBand")
    add_band(
        slide,
        x,
        header_y + header_h + row_h,
        w,
        row_h,
        fill=PURPLE_TINT,
        name="RESULTS_TableBestBand",
    )

    columns = [
        (x, col_model, "Model", PP_ALIGN.LEFT),
        (x + col_model, col_recall, "Recall@100", PP_ALIGN.RIGHT),
        (x + col_model + col_recall, col_delta, "Δ", PP_ALIGN.RIGHT),
    ]
    for index, (cell_x, cell_w, label, align) in enumerate(columns):
        add_textbox(
            slide,
            cell_x,
            header_y,
            cell_w,
            header_h,
            label,
            size=30,
            bold=True,
            color=MUTED,
            align=align,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"RESULTS_TableHeader_{index}",
            margins=0.10,
        )

    rows = [
        ("SASRec-Audio-Uniform", "0.12303", "baseline", False),
        ("RQ-VAE + item LogQ", "0.14138", "+14.9%", True),
    ]
    for row_index, (model, recall, delta, is_best) in enumerate(rows):
        y = header_y + header_h + row_index * row_h
        values = [
            (x, col_model, model, PP_ALIGN.LEFT),
            (x + col_model, col_recall, recall, PP_ALIGN.RIGHT),
            (x + col_model + col_recall, col_delta, delta, PP_ALIGN.RIGHT),
        ]
        for col_index, (cell_x, cell_w, value, align) in enumerate(values):
            add_textbox(
                slide,
                cell_x,
                y,
                cell_w,
                row_h,
                value,
                size=36 if is_best and col_index > 0 else 30,
                bold=is_best,
                color=GREEN_DARK if is_best and col_index > 0 else INK,
                align=align,
                valign=MSO_VERTICAL_ANCHOR.MIDDLE,
                name=f"RESULTS_TableCell_{row_index}_{col_index}",
                margins=0.10,
            )
        add_rule(
            slide,
            x,
            y + row_h,
            x + w,
            y + row_h,
            color=PURPLE_RULE,
            width=1,
            name=f"RESULTS_TableRule_{row_index}",
        )

    add_textbox(
        slide,
        x,
        22.02,
        w,
        0.62,
        "Best SID: RQ-VAE + item LogQ, L=2, C=256",
        size=30,
        bold=True,
        color=GREEN_DARK,
        name="RESULTS_GenerationTableNote",
    )


def build_results(slide, heatmap) -> None:
    x, w = 31.19, 13.78
    add_textbox(
        slide,
        x,
        9.50,
        w,
        1.02,
        "КАКИЕ ХАРАКТЕРИСТИКИ SID СВЯЗАНЫ\nС DOWNSTREAM-КАЧЕСТВОМ?",
        size=34,
        bold=True,
        name="RESULTS_HeatmapQuestion",
        line_spacing=1.0,
    )

    heatmap.left = Inches(30.84)
    heatmap.top = Inches(10.72)
    heatmap.width = Inches(14.48)
    heatmap.height = Inches(5.53)

    add_inline_statement(
        slide,
        x,
        16.42,
        w,
        0.72,
        "DCNv2 Pairwise",
        "ICR и code entropy: r = 0.47.",
        name="RESULTS_Observation_1",
        label_color=GREEN_DARK,
    )
    add_inline_statement(
        slide,
        x,
        17.16,
        w,
        0.72,
        "DCNv2 ROC-AUC",
        "нет общей зависимости; max |r| = 0.20.",
        name="RESULTS_Observation_2",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        18.04,
        x + w,
        18.04,
        color=PURPLE_RULE,
        width=1.5,
        name="RESULTS_ObservationRule",
    )
    build_generation_table(slide)


def build_conclusions(slide) -> None:
    x_positions = [1.96, 12.63, 23.30, 33.97]
    col_w = 9.85
    for index in (1, 2, 3):
        add_rule(
            slide,
            x_positions[index] - 0.40,
            28.90,
            x_positions[index] - 0.40,
            31.02,
            color=PURPLE_RULE,
            width=1.5,
            name=f"CONCLUSIONS_Divider_{index}",
        )

    theses = [
        (
            "SID ≠ downstream quality",
            "Хорошие внутренние характеристики SID не гарантируют улучшение итоговой модели.",
        ),
        (
            "Ranking: устойчивого улучшения нет",
            "SID-конфигурации не превзошли item-only baseline.",
        ),
        (
            "Generation: +14.9%",
            "RQ-VAE + item LogQ дал лучший Recall@100 относительно baseline.",
        ),
        (
            "Проверять нужно downstream",
            "SID нужно оценивать через реальные downstream-задачи, а не только внутренние метрики.",
        ),
    ]
    for index, ((head, body), x) in enumerate(zip(theses, x_positions), start=1):
        add_rich_textbox(
            slide,
            x,
            28.82,
            col_w,
            2.30,
            [
                {"runs": [(head, 34, True, GREEN_DARK)], "line_spacing": 1.0},
                {"runs": [(body, 30, False, INK)], "before": 7, "line_spacing": 1.02},
            ],
            name=f"CONCLUSIONS_Thesis_{index}",
        )


def assert_invariants(prs, *, rqvae_hash: str, heatmap_hash: str) -> None:
    assert len(prs.slides) == 1
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    shapes = {shape.name: shape for shape in slide.shapes}
    required = {
        "INTRO_RQKMeansBand",
        "INTRO_ResearchQuestion",
        "METHODS_BaselineBand",
        "METHODS_UserAwareBand",
        "METHODS_PSIBand",
        "METHODS_R3VAEBand",
        "METHODS_LogQBand",
        "Рисунок 4",
        "RESULTS_Heatmap_Full",
        "RESULTS_TableBestBand",
        "CONCLUSIONS_Thesis_4",
    }
    assert required <= shapes.keys(), f"Missing shapes: {sorted(required - shapes.keys())}"
    assert "RESULTS_Recall_vs_ICR_SVG" not in shapes
    assert sha256(shapes["Рисунок 4"].image.blob).hexdigest() == rqvae_hash
    assert sha256(shapes["RESULTS_Heatmap_Full"].image.blob).hexdigest() == heatmap_hash

    all_text = "\n".join(
        shape.text for shape in slide.shapes if getattr(shape, "has_text_frame", False)
    )
    assert "Ключевые SID-метрики" not in all_text
    assert "Independent Code Rate (ICR)" not in all_text

    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"
        if not shape.name.startswith(("INTRO_", "METHODS_", "RESULTS_", "CONCLUSIONS_")):
            continue
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        assert run.font.size.pt >= 30, (
                            f"Text below 30 pt in {shape.name}: {run.text!r} ({run.font.size.pt})"
                        )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    slide = prs.slides[0]
    original_shapes = {shape.name: shape for shape in slide.shapes}
    rqvae_picture = original_shapes["Рисунок 4"]
    heatmap = original_shapes["RESULTS_Heatmap_Full"]
    rqvae_hash = sha256(rqvae_picture.image.blob).hexdigest()
    heatmap_hash = sha256(heatmap.image.blob).hexdigest()

    delete_named_content(
        slide,
        ("INTRO_", "METHODS_", "RESULTS_", "CONCLUSIONS_"),
        keep={"RESULTS_Heatmap_Full"},
    )
    delete_hidden_named_content(slide, ("INTRO_", "METHODS_", "RESULTS_", "CONCLUSIONS_"))

    build_introduction(slide)
    build_methods(slide, rqvae_picture)
    build_results(slide, heatmap)
    build_conclusions(slide)

    prs.core_properties.comments = (
        "Hallmark · pre-emit P5 H5 E4 S5 R5 V4 · constrained Narrative Workflow · "
        "poster_0 source preserved · "
        "original RQ-VAE and heatmap images retained byte-for-byte"
    )
    assert_invariants(prs, rqvae_hash=rqvae_hash, heatmap_hash=heatmap_hash)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
