from __future__ import annotations

# Constrained poster update · current spaced poster is the sole base

from hashlib import sha256
from pathlib import Path
import tempfile
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from refine_poster_0_spacing import EMU_PER_INCH, set_box, shape_map
from update_poster_results import (
    inject_native_svg_layers,
    render_svg_fallbacks,
    set_picture_metadata,
)


SOURCE = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_0_improved_spaced.pptx"
)
RECALL_SVG = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "reports/figures/all_methods_sid_quality/recall100/recall100_vs_icr.svg"
)
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_0_improved_spaced_with_recall_plot.pptx"
)

RESULTS_X = 31.19
RESULTS_W = 13.78
PLOT_NAME = "RESULTS_Recall_vs_ICR_SVG"


def reposition_results(shapes) -> None:
    # Keep the heatmap broad enough for its long metric labels.
    set_box(shapes["RESULTS_Heatmap_Full"], 31.535, 10.65, 13.09, 5.00)

    set_box(shapes["RESULTS_Observation_1"], RESULTS_X, 15.82, RESULTS_W, 0.56)
    set_box(shapes["RESULTS_Observation_2"], RESULTS_X, 16.49, RESULTS_W, 0.56)
    set_box(shapes["RESULTS_ObservationRule"], RESULTS_X, 17.18, RESULTS_W, 0)

    set_box(shapes["RESULTS_GenerationTableHead"], RESULTS_X, 17.52, RESULTS_W, 0.50)
    header_y, header_h, row_h = 18.10, 0.66, 0.82
    set_box(shapes["RESULTS_TableHeaderBand"], RESULTS_X, header_y, RESULTS_W, header_h)
    set_box(
        shapes["RESULTS_TableBestBand"],
        RESULTS_X,
        header_y + header_h + row_h,
        RESULTS_W,
        row_h,
    )

    columns = [
        (RESULTS_X, 7.25),
        (RESULTS_X + 7.25, 3.35),
        (RESULTS_X + 10.60, 3.18),
    ]
    for index, (cell_x, cell_w) in enumerate(columns):
        cell = shapes[f"RESULTS_TableHeader_{index}"]
        set_box(cell, cell_x, header_y, cell_w, header_h)
        cell.text_frame.margin_top = Inches(0.10)
        cell.text_frame.margin_bottom = Inches(0.10)

    for row_index in range(2):
        row_y = header_y + header_h + row_index * row_h
        for col_index, (cell_x, cell_w) in enumerate(columns):
            cell = shapes[f"RESULTS_TableCell_{row_index}_{col_index}"]
            set_box(cell, cell_x, row_y, cell_w, row_h)
            cell.text_frame.margin_top = Inches(0.12)
            cell.text_frame.margin_bottom = Inches(0.12)
        set_box(shapes[f"RESULTS_TableRule_{row_index}"], RESULTS_X, row_y + row_h, RESULTS_W, 0)

    set_box(shapes["RESULTS_GenerationTableNote"], RESULTS_X, 20.58, RESULTS_W, 0.50)


def add_recall_plot(slide, fallback_png: Path) -> None:
    width = 11.75
    left = RESULTS_X + (RESULTS_W - width) / 2
    shape = slide.shapes.add_picture(
        str(fallback_png),
        Inches(left),
        Inches(21.16),
        width=Inches(width),
    )
    set_picture_metadata(
        shape,
        name=PLOT_NAME,
        description=(
            "Recall@100 versus raw independent code rate (ICR), with panels for "
            "varying number of codebooks and varying codebook size."
        ),
    )
    assert shape.top + shape.height <= Inches(26.40)


def assert_invariants(prs: Presentation, *, rqvae_hash: str, heatmap_hash: str) -> None:
    assert len(prs.slides) == 1
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    shapes = shape_map(slide)
    assert PLOT_NAME in shapes
    assert sha256(shapes["Рисунок 4"].image.blob).hexdigest() == rqvae_hash
    assert sha256(shapes["RESULTS_Heatmap_Full"].image.blob).hexdigest() == heatmap_hash

    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"
        if not shape.name.startswith(("INTRO_", "METHODS_", "RESULTS_", "CONCLUSIONS_")):
            continue
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        assert run.font.size.pt >= 30, (
                            f"Text below 30 pt in {shape.name}: {run.text!r}"
                        )


def main() -> None:
    for path in (SOURCE, RECALL_SVG):
        if not path.exists():
            raise FileNotFoundError(path)

    with tempfile.TemporaryDirectory(prefix="poster-recall-plot-") as temporary:
        temp_dir = Path(temporary)
        fallback_png = render_svg_fallbacks([RECALL_SVG], temp_dir)[RECALL_SVG]
        with Image.open(fallback_png) as image:
            assert image.width / image.height > 2.1

        prs = Presentation(SOURCE)
        slide = prs.slides[0]
        shapes = shape_map(slide)
        rqvae_hash = sha256(shapes["Рисунок 4"].image.blob).hexdigest()
        heatmap_hash = sha256(shapes["RESULTS_Heatmap_Full"].image.blob).hexdigest()

        reposition_results(shapes)
        add_recall_plot(slide, fallback_png)
        prs.core_properties.comments = (
            "Current spaced poster preserved · recall100_vs_icr added to Results · "
            "native SVG with PNG fallback · A0 · supplied RQ-VAE and heatmap retained"
        )
        assert_invariants(prs, rqvae_hash=rqvae_hash, heatmap_hash=heatmap_hash)

        base_pptx = temp_dir / "poster-base.pptx"
        prs.save(base_pptx)
        inject_native_svg_layers(
            base_pptx,
            OUTPUT,
            [(PLOT_NAME, RECALL_SVG, "poster_results_recall_vs_icr.svg")],
        )

    with zipfile.ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
        assert "ppt/media/poster_results_recall_vs_icr.svg" in archive.namelist()

    reopened = Presentation(OUTPUT)
    assert_invariants(reopened, rqvae_hash=rqvae_hash, heatmap_hash=heatmap_hash)
    print(OUTPUT)


if __name__ == "__main__":
    main()
