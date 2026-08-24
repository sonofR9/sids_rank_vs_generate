from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_poster_pptx import (
    GREEN,
    GREEN_DARK,
    INK,
    MUTED,
    add_rich_textbox,
    add_rule,
    add_textbox,
)


SOURCE = Path("/Users/konstantinryadinskiy/Desktop/poster_semantic_id_final_equations.pptx")
HEATMAP = Path("/Users/konstantinryadinskiy/Downloads/heat_map_deduplicated_horizontal.png")
RECALL_SVG = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "reports/figures/all_methods_sid_quality/recall100/recall100_vs_icr.svg"
)
TRANSACT_SVG = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "reports/figures/all_methods_sid_quality/transact/"
    "transact_vs_intra_code_similarity.svg"
)
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_semantic_id_final_visual_results.pptx"
)

SOFFICE_FALLBACK = Path(
    "/Users/konstantinryadinskiy/.cache/codex-runtimes/codex-primary-runtime/"
    "dependencies/bin/override/soffice"
)

EMU_PER_INCH = 914400
RESULTS_X = 31.19
RESULTS_W = 13.78

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
ASVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"
IMAGE_REL_TYPE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
)


def delete_named_results_shapes(slide) -> None:
    """Remove only the generated result-layer shapes; preserve the template card and title."""
    for shape in list(slide.shapes):
        if shape.name.startswith("RESULTS_"):
            element = shape._element
            element.getparent().remove(element)


def add_generation_highlight(slide) -> None:
    """Keep the downstream win as the first and strongest result in the column."""
    add_rule(
        slide,
        RESULTS_X,
        9.47,
        RESULTS_X + RESULTS_W,
        9.47,
        color=GREEN,
        width=3,
        name="RESULTS_GenerationRule",
    )
    add_textbox(
        slide,
        RESULTS_X,
        9.54,
        RESULTS_W,
        0.40,
        "GENERATION УЛУЧШИЛСЯ · Recall@100",
        size=32,
        bold=True,
        color=GREEN_DARK,
        name="RESULTS_GenerationHead",
        line_spacing=1.0,
    )
    add_textbox(
        slide,
        RESULTS_X,
        9.94,
        8.15,
        0.65,
        "0.12303 → 0.14138",
        size=50,
        bold=True,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="RESULTS_GenerationValues",
        line_spacing=1.0,
    )
    add_textbox(
        slide,
        RESULTS_X + 8.23,
        9.90,
        5.47,
        0.72,
        "+14.9%",
        size=60,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.RIGHT,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="RESULTS_GenerationDelta",
        line_spacing=1.0,
    )
    add_rich_textbox(
        slide,
        RESULTS_X,
        10.55,
        RESULTS_W,
        0.47,
        [
            {
                "runs": [
                    ("Baseline · ", 30, True, MUTED),
                    ("SASRec-Audio-Uniform", 30, False, INK),
                ],
                "line_spacing": 1.0,
            }
        ],
        name="RESULTS_GenerationBaseline",
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )
    add_rich_textbox(
        slide,
        RESULTS_X,
        11.00,
        RESULTS_W,
        0.47,
        [
            {
                "runs": [
                    ("Best SID · ", 30, True, MUTED),
                    ("RQ-VAE + item LogQ, L=2, C=256", 30, False, INK),
                ],
                "line_spacing": 1.0,
            }
        ],
        name="RESULTS_GenerationBest",
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )


def crop_external_whitespace(source: Path, target: Path, pad_px: int = 12) -> tuple[int, int]:
    """Losslessly crop only the near-white outer margin of the supplied PNG."""
    with Image.open(source) as image:
        rgba = image.convert("RGBA")
        background = Image.new("RGBA", rgba.size, rgba.getpixel((0, 0)))
        difference = ImageChops.difference(rgba, background).convert("RGB")
        mask = difference.convert("L").point(lambda value: 255 if value > 3 else 0)
        bbox = mask.getbbox()
        if bbox is None:
            bbox = (0, 0, rgba.width, rgba.height)
        left = max(0, bbox[0] - pad_px)
        top = max(0, bbox[1] - pad_px)
        right = min(rgba.width, bbox[2] + pad_px)
        bottom = min(rgba.height, bbox[3] + pad_px)
        cropped = rgba.crop((left, top, right, bottom))
        cropped.save(target, format="PNG", optimize=False)
        return cropped.size


def set_picture_metadata(shape, *, name: str, description: str) -> None:
    shape.name = name
    c_nv_pr = shape._element.nvPicPr.cNvPr
    c_nv_pr.set("descr", description)


def add_heatmap(slide, cropped_png: Path, pixel_size: tuple[int, int]) -> None:
    width = RESULTS_W
    height = width * pixel_size[1] / pixel_size[0]
    shape = slide.shapes.add_picture(
        str(cropped_png),
        Inches(RESULTS_X),
        Inches(11.58),
        width=Inches(width),
        height=Inches(height),
    )
    set_picture_metadata(
        shape,
        name="RESULTS_Heatmap_Full",
        description=(
            "Pearson correlation heatmap: all supplied SID metrics versus downstream "
            "metrics, including title, subtitle, cell values, and Pearson r colorbar."
        ),
    )


def render_svg_fallbacks(svg_paths: list[Path], target_dir: Path) -> dict[Path, Path]:
    soffice = shutil.which("soffice")
    if soffice is None and SOFFICE_FALLBACK.exists():
        soffice = str(SOFFICE_FALLBACK)
    if soffice is None:
        raise RuntimeError("LibreOffice/soffice is required to create PNG SVG fallbacks")

    profile_dir = target_dir / "lo-profile"
    command = [
        soffice,
        f"-env:UserInstallation={profile_dir.as_uri()}",
        "--headless",
        "--convert-to",
        "png",
        "--outdir",
        str(target_dir),
        *[str(path) for path in svg_paths],
    ]
    completed = subprocess.run(command, check=True, capture_output=True, text=True)
    fallbacks: dict[Path, Path] = {}
    for svg_path in svg_paths:
        png_path = target_dir / f"{svg_path.stem}.png"
        if not png_path.exists():
            raise RuntimeError(
                f"Fallback render missing for {svg_path.name}.\n"
                f"stdout: {completed.stdout}\nstderr: {completed.stderr}"
            )
        fallbacks[svg_path] = png_path
    return fallbacks


def add_svg_fallback_picture(
    slide,
    fallback_png: Path,
    *,
    name: str,
    description: str,
    y: float,
) -> None:
    width = 10.45
    left = RESULTS_X + (RESULTS_W - width) / 2
    shape = slide.shapes.add_picture(
        str(fallback_png),
        Inches(left),
        Inches(y),
        width=Inches(width),
    )
    set_picture_metadata(shape, name=name, description=description)


def _next_relationship_id(rels_root: etree._Element) -> str:
    values = []
    for relationship in rels_root:
        match = re.fullmatch(r"rId(\d+)", relationship.get("Id", ""))
        if match:
            values.append(int(match.group(1)))
    return f"rId{max(values, default=0) + 1}"


def inject_native_svg_layers(
    base_pptx: Path,
    output_pptx: Path,
    mappings: list[tuple[str, Path, str]],
) -> None:
    """Add Office SVG extensions while retaining PNG fallbacks for compatibility."""
    slide_xml_name = "ppt/slides/slide1.xml"
    slide_rels_name = "ppt/slides/_rels/slide1.xml.rels"
    content_types_name = "[Content_Types].xml"

    with zipfile.ZipFile(base_pptx, "r") as source_zip:
        slide_root = etree.fromstring(source_zip.read(slide_xml_name))
        rels_root = etree.fromstring(source_zip.read(slide_rels_name))
        content_types_root = etree.fromstring(source_zip.read(content_types_name))

        changed_entries: dict[str, bytes] = {}
        for picture_name, svg_path, media_name in mappings:
            pictures = slide_root.xpath(
                ".//p:pic[p:nvPicPr/p:cNvPr[@name=$picture_name]]",
                namespaces={"p": P_NS, "a": A_NS},
                picture_name=picture_name,
            )
            if len(pictures) != 1:
                raise AssertionError(
                    f"Expected one fallback picture named {picture_name!r}, found {len(pictures)}"
                )
            blip = pictures[0].find(f".//{{{A_NS}}}blip")
            if blip is None:
                raise AssertionError(f"Picture {picture_name!r} has no DrawingML blip")

            relationship_id = _next_relationship_id(rels_root)
            relationship = etree.SubElement(
                rels_root, etree.QName(REL_NS, "Relationship")
            )
            relationship.set("Id", relationship_id)
            relationship.set("Type", IMAGE_REL_TYPE)
            relationship.set("Target", f"../media/{media_name}")

            ext_list = blip.find(f"{{{A_NS}}}extLst")
            if ext_list is None:
                ext_list = etree.SubElement(blip, etree.QName(A_NS, "extLst"))
            extension = etree.SubElement(ext_list, etree.QName(A_NS, "ext"))
            extension.set("uri", SVG_EXT_URI)
            svg_blip = etree.SubElement(
                extension,
                etree.QName(ASVG_NS, "svgBlip"),
                nsmap={"asvg": ASVG_NS},
            )
            svg_blip.set(etree.QName(R_NS, "embed"), relationship_id)
            changed_entries[f"ppt/media/{media_name}"] = svg_path.read_bytes()

        svg_defaults = content_types_root.xpath(
            "/ct:Types/ct:Default[@Extension='svg']",
            namespaces={"ct": CT_NS},
        )
        if not svg_defaults:
            default = etree.SubElement(
                content_types_root, etree.QName(CT_NS, "Default")
            )
            default.set("Extension", "svg")
            default.set("ContentType", "image/svg+xml")

        changed_entries[slide_xml_name] = etree.tostring(
            slide_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        changed_entries[slide_rels_name] = etree.tostring(
            rels_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )
        changed_entries[content_types_name] = etree.tostring(
            content_types_root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        with zipfile.ZipFile(output_pptx, "w", zipfile.ZIP_DEFLATED) as target_zip:
            for item in source_zip.infolist():
                if item.filename in changed_entries:
                    continue
                target_zip.writestr(item, source_zip.read(item.filename))
            for name, content in changed_entries.items():
                target_zip.writestr(name, content)


def assert_poster_invariants(prs: Presentation) -> None:
    assert len(prs.slides) == 1, "Poster must remain a one-slide deck"
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"

    names = {shape.name for shape in slide.shapes}
    required = {
        "RESULTS_GenerationHead",
        "RESULTS_GenerationValues",
        "RESULTS_GenerationDelta",
        "RESULTS_Heatmap_Full",
        "RESULTS_Recall_vs_ICR_SVG",
        "RESULTS_TransAct_vs_IntraCode_SVG",
    }
    assert required <= names, f"Missing result shapes: {sorted(required - names)}"
    assert not any(name.startswith("RESULTS_DCNv2") for name in names)
    assert not any(name.startswith("RESULTS_TransActRule") for name in names)
    assert not any(name.startswith("RESULTS_TransActHead") for name in names)
    assert not any(name.startswith("RESULTS_TransActValues") for name in names)
    assert not any(name.startswith("RESULTS_TransActDelta") for name in names)
    assert not any(name.startswith("RESULTS_TransActModels") for name in names)

    math_namespace = "http://schemas.microsoft.com/office/drawing/2010/main"
    office_math_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    math_objects = list(slide._element.iter(f"{{{math_namespace}}}m"))
    fractions = list(slide._element.iter(f"{{{office_math_namespace}}}f"))
    subscripts = list(slide._element.iter(f"{{{office_math_namespace}}}sSub"))
    assert len(math_objects) == 3, "All three editable Office Math objects must be preserved"
    assert len(fractions) == 2, "Both Office Math fractions must be preserved"
    assert len(subscripts) == 4, "All four Office Math subscripts must be preserved"


def main() -> None:
    for path in (SOURCE, HEATMAP, RECALL_SVG, TRANSACT_SVG):
        if not path.exists():
            raise FileNotFoundError(path)

    with tempfile.TemporaryDirectory(prefix="poster-results-") as temporary:
        temp_dir = Path(temporary)
        cropped_heatmap = temp_dir / "heatmap-cropped.png"
        pixel_size = crop_external_whitespace(HEATMAP, cropped_heatmap)
        fallbacks = render_svg_fallbacks([RECALL_SVG, TRANSACT_SVG], temp_dir)

        prs = Presentation(SOURCE)
        slide = prs.slides[0]
        delete_named_results_shapes(slide)
        add_generation_highlight(slide)
        add_heatmap(slide, cropped_heatmap, pixel_size)
        add_svg_fallback_picture(
            slide,
            fallbacks[RECALL_SVG],
            name="RESULTS_Recall_vs_ICR_SVG",
            description=(
                "Recall@100 versus raw independent code rate (ICR), with panels for "
                "varying number of codebooks and varying codebook size."
            ),
            y=16.93,
        )
        add_svg_fallback_picture(
            slide,
            fallbacks[TRANSACT_SVG],
            name="RESULTS_TransAct_vs_IntraCode_SVG",
            description=(
                "TransAct pairwise accuracy versus intra-code similarity, with panels "
                "for varying number of codebooks and varying codebook size."
            ),
            y=21.65,
        )

        prs.core_properties.comments = (
            "Hallmark · constrained fixed-template poster update · technical/austere · "
            "generation result first · supplied full heatmap · native SVG charts with "
            "PNG fallbacks · pre-emit critique: P5 H5 E5 S5 R5 V4"
        )
        assert_poster_invariants(prs)

        base_pptx = temp_dir / "poster-base.pptx"
        prs.save(base_pptx)
        inject_native_svg_layers(
            base_pptx,
            OUTPUT,
            [
                (
                    "RESULTS_Recall_vs_ICR_SVG",
                    RECALL_SVG,
                    "poster_results_recall_vs_icr.svg",
                ),
                (
                    "RESULTS_TransAct_vs_IntraCode_SVG",
                    TRANSACT_SVG,
                    "poster_results_transact_vs_intra.svg",
                ),
            ],
        )

    print(OUTPUT)


if __name__ == "__main__":
    main()
