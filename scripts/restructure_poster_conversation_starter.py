from __future__ import annotations

# Hallmark · pre-emit critique: P5 H5 E4 S5 R5 V5
# Narrative Workflow · technical/austere · fixed-template A0 scientific poster

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_poster_pptx import (
    GREEN,
    GREEN_DARK,
    INK,
    MUTED,
    PURPLE_RULE,
    PURPLE_TINT,
    add_band,
    add_flow_arrow,
    add_rich_textbox,
    add_rule,
    add_textbox,
)


SOURCE = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_semantic_id_final_visual_results_polished.pptx"
)
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_semantic_id_conversation_starter.pptx"
)

EMU_PER_INCH = 914400
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def delete_named_content(slide, prefixes: tuple[str, ...], *, keep: set[str] | None = None) -> None:
    """Remove the old content layer while preserving the fixed ML-Академия template."""
    keep = keep or set()
    for shape in list(slide.shapes):
        if shape.name in keep:
            continue
        if shape.name.startswith(prefixes):
            shape._element.getparent().remove(shape._element)


def delete_hidden_named_content(slide, prefixes: tuple[str, ...]) -> None:
    """Remove AlternateContent wrappers that python-pptx does not expose as shapes."""
    c_nv_pr_tag = f"{{{P_NS}}}cNvPr"
    alternate_tag = f"{{{MC_NS}}}AlternateContent"
    wrappers = set()
    for node in slide._element.iter(c_nv_pr_tag):
        if not node.get("name", "").startswith(prefixes):
            continue
        parent = node.getparent()
        while parent is not None and parent.tag != alternate_tag:
            parent = parent.getparent()
        if parent is not None:
            wrappers.add(parent)
    for wrapper in wrappers:
        wrapper.getparent().remove(wrapper)


def delete_residual_center_content(slide) -> None:
    """Clear legacy artwork inside the methods card, but keep the fixed card and heading."""
    left, top, right, bottom = (15.40, 9.20, 29.95, 26.65)
    for shape in list(slide.shapes):
        x = shape.left / EMU_PER_INCH
        y = shape.top / EMU_PER_INCH
        x2 = (shape.left + shape.width) / EMU_PER_INCH
        y2 = (shape.top + shape.height) / EMU_PER_INCH
        if x >= left and y >= top and x2 <= right and y2 <= bottom:
            shape._element.getparent().remove(shape._element)


def add_inline_statement(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    body: str,
    *,
    name: str,
    label_color=MUTED,
    body_color=INK,
    size: float = 30,
):
    return add_rich_textbox(
        slide,
        x,
        y,
        w,
        h,
        [
            {
                "runs": [
                    (f"{label} · ", size, True, label_color),
                    (body, size, False, body_color),
                ],
                "line_spacing": 1.02,
            }
        ],
        name=name,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )


def build_introduction(slide) -> None:
    x, w = 1.86, 12.79

    add_textbox(
        slide,
        x,
        9.50,
        w,
        0.48,
        "Зачем нужны Semantic IDs?",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_WhyHead",
    )
    add_textbox(
        slide,
        x,
        10.04,
        w,
        0.92,
        "Semantic IDs — дискретные коды объектов для генеративных рекомендательных моделей.",
        size=32,
        bold=True,
        name="INTRO_WhyBody",
        line_spacing=1.02,
    )
    add_rule(
        slide,
        x,
        11.18,
        x + w,
        11.18,
        color=PURPLE_RULE,
        width=1.5,
        name="INTRO_WhyRule",
    )

    add_textbox(
        slide,
        x,
        11.42,
        w,
        0.48,
        "Как получается SID?",
        size=34,
        bold=True,
        name="INTRO_HowHead",
    )
    add_textbox(
        slide,
        x,
        11.94,
        w,
        0.52,
        "RQ-KMeans · последовательное residual quantization",
        size=30,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_RQKMeansLabel",
    )
    add_band(slide, x, 12.52, w, 2.22, name="INTRO_RQKMeansBand")

    stage_y = 12.76
    stage_h = 0.70
    add_textbox(
        slide,
        x + 0.12,
        stage_y,
        2.55,
        stage_h,
        "item embedding",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansStage_Embedding",
    )
    add_flow_arrow(slide, x + 2.70, stage_y + 0.08, w=0.36, name="INTRO_RQKMeansArrow_1")
    add_textbox(
        slide,
        x + 3.08,
        stage_y,
        2.10,
        stage_h,
        "level 1 · c₁",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansStage_Level1",
    )
    add_flow_arrow(slide, x + 5.22, stage_y + 0.08, w=0.36, name="INTRO_RQKMeansArrow_2")
    add_textbox(
        slide,
        x + 5.62,
        stage_y,
        2.05,
        stage_h,
        "residual",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansStage_Residual",
    )
    add_flow_arrow(slide, x + 7.72, stage_y + 0.08, w=0.36, name="INTRO_RQKMeansArrow_3")
    add_textbox(
        slide,
        x + 8.12,
        stage_y,
        2.52,
        stage_h,
        "levels 2…L",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansStage_Next",
    )
    add_flow_arrow(slide, x + 10.66, stage_y + 0.08, w=0.36, name="INTRO_RQKMeansArrow_4")
    add_textbox(
        slide,
        x + 11.03,
        stage_y,
        1.62,
        stage_h,
        "SID",
        size=32,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansStage_SID",
    )
    add_textbox(
        slide,
        x + 0.20,
        13.77,
        w - 0.40,
        0.66,
        "c₁  →  c₂  →  …  →  cₗ   =   Semantic ID",
        size=36,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="INTRO_RQKMeansSequence",
    )
    add_textbox(
        slide,
        x,
        14.91,
        w,
        0.52,
        "RQ-VAE — обучаемый baseline с той же residual-логикой.",
        size=30,
        color=MUTED,
        name="INTRO_RQVAEReference",
    )

    add_textbox(
        slide,
        x,
        15.58,
        w,
        0.48,
        "Что хотим выяснить?",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_QuestionHead",
    )
    add_rich_textbox(
        slide,
        x,
        16.12,
        w,
        2.32,
        [
            {
                "runs": [
                    (
                        "Какие свойства и способы построения SID действительно "
                        "отражаются на downstream-качестве?",
                        40,
                        True,
                        INK,
                    )
                ],
                "line_spacing": 1.0,
            }
        ],
        name="INTRO_ResearchQuestion",
    )
    add_rule(
        slide,
        x,
        18.52,
        x + 7.30,
        18.52,
        color=GREEN,
        width=4,
        name="INTRO_QuestionUnderline",
    )

    add_textbox(
        slide,
        x,
        18.79,
        w,
        0.48,
        "Где проверяем?",
        size=34,
        bold=True,
        name="INTRO_EvaluationHead",
    )
    add_textbox(
        slide,
        x,
        19.29,
        w,
        0.46,
        "Yambda-50M · time split",
        size=30,
        bold=True,
        color=MUTED,
        name="INTRO_Dataset",
    )
    add_rule(
        slide,
        x + 6.22,
        19.94,
        x + 6.22,
        23.82,
        color=PURPLE_RULE,
        width=1.5,
        name="INTRO_TaskDivider",
    )
    add_textbox(
        slide,
        x,
        19.94,
        5.92,
        0.48,
        "Генерация",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_GenerationHead",
    )
    add_rich_textbox(
        slide,
        x,
        20.50,
        5.92,
        3.20,
        [
            {"runs": [("TIGER → SID следующего объекта", 30, True, INK)]},
            {"runs": [("Baseline · SASRec по item ID", 30, False, INK)], "before": 6},
            {"runs": [("Метрика · Recall@100", 30, False, INK)], "before": 6},
        ],
        name="INTRO_GenerationBody",
    )
    add_textbox(
        slide,
        x + 6.52,
        19.94,
        6.27,
        0.48,
        "Ранжирование",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_RankingHead",
    )
    add_rich_textbox(
        slide,
        x + 6.52,
        20.50,
        6.27,
        3.38,
        [
            {"runs": [("DCNv2 + SID целевого объекта", 30, True, INK)]},
            {"runs": [("SID — дополнительный признак", 30, False, INK)], "before": 6},
            {
                "runs": [("Метрики · ROC-AUC / Pairwise Accuracy", 30, False, INK)],
                "before": 6,
                "line_spacing": 1.02,
            },
        ],
        name="INTRO_RankingBody",
    )


def build_methods(slide) -> None:
    x, w = 15.78, 13.78

    add_textbox(
        slide,
        x,
        9.50,
        w,
        0.54,
        "Что именно мы меняли в SID?",
        size=36,
        bold=True,
        name="METHODS_QuestionHead",
    )
    add_band(slide, x, 10.15, w, 1.18, name="METHODS_RQVAEReferenceBand")
    add_textbox(
        slide,
        x + 0.18,
        10.34,
        2.05,
        0.78,
        "RQ-VAE",
        size=34,
        bold=True,
        color=GREEN_DARK,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_RQVAEReferenceHead",
    )
    add_textbox(
        slide,
        x + 2.32,
        10.30,
        w - 2.50,
        0.84,
        "Reference point · residual quantization по обучаемым codebooks.",
        size=30,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_RQVAEReferenceBody",
    )

    add_textbox(
        slide,
        x,
        11.66,
        w,
        0.56,
        "User-aware",
        size=36,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_UserAwareHead",
    )
    add_inline_statement(
        slide,
        x,
        12.40,
        w,
        0.68,
        "Что меняем",
        "добавляем информацию о пользовательских предпочтениях.",
        name="METHODS_UserAwareWhat",
    )
    add_band(slide, x, 13.22, w, 1.10, name="METHODS_UserAwareBand")
    add_textbox(
        slide,
        x + 0.30,
        13.39,
        2.18,
        0.72,
        "ITEMS",
        size=32,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_UserAwareItems",
    )
    add_textbox(
        slide,
        x + 2.52,
        13.39,
        0.55,
        0.72,
        "+",
        size=34,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_UserAwarePlus",
    )
    add_textbox(
        slide,
        x + 3.12,
        13.39,
        3.38,
        0.72,
        "USER SIGNALS",
        size=32,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_UserAwareSignals",
    )
    add_flow_arrow(slide, x + 6.62, 13.46, w=0.50, name="METHODS_UserAwareArrow")
    add_textbox(
        slide,
        x + 7.22,
        13.39,
        6.15,
        0.72,
        "USER-AWARE SID",
        size=34,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_UserAwareSID",
    )
    add_inline_statement(
        slide,
        x,
        14.42,
        w,
        0.68,
        "Зачем",
        "согласовать коды с пользовательским восприятием объектов.",
        name="METHODS_UserAwareWhy",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        15.20,
        x + w,
        15.20,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_UserAwareRule",
    )

    add_textbox(
        slide,
        x,
        15.46,
        w,
        0.56,
        "PSI",
        size=36,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_PSIHead",
    )
    add_inline_statement(
        slide,
        x,
        16.18,
        w,
        0.68,
        "Что меняем",
        "при коллизии выбираем следующую близкую свободную комбинацию.",
        name="METHODS_PSIWhat",
    )
    add_band(slide, x, 17.00, w, 1.12, name="METHODS_PSIBand")
    add_textbox(
        slide,
        x + 0.20,
        17.18,
        3.05,
        0.72,
        "nearest SID",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_PSINearest",
    )
    add_flow_arrow(slide, x + 3.30, 17.25, w=0.42, name="METHODS_PSIArrow_1")
    add_textbox(
        slide,
        x + 3.78,
        17.18,
        3.20,
        0.72,
        "occupied  ×",
        size=30,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_PSIOccupied",
    )
    add_flow_arrow(slide, x + 7.04, 17.25, w=0.42, name="METHODS_PSIArrow_2")
    add_textbox(
        slide,
        x + 7.52,
        17.18,
        6.02,
        0.72,
        "next nearest free SID  ✓",
        size=30,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_PSIFree",
    )
    add_inline_statement(
        slide,
        x,
        18.24,
        w,
        0.68,
        "Зачем",
        "устранить коллизии без дополнительного токена.",
        name="METHODS_PSIWhy",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        19.02,
        x + w,
        19.02,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_PSIRule",
    )

    left_w, gap = 6.58, 0.62
    right_x = x + left_w + gap
    add_textbox(
        slide,
        x,
        19.32,
        left_w,
        0.56,
        "R3-VAE",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_R3VAEHead",
    )
    add_inline_statement(
        slide,
        x,
        20.04,
        left_w,
        1.20,
        "Что меняем",
        "обучаемый semantic anchor для первого residual.",
        name="METHODS_R3VAEWhat",
    )
    add_band(slide, x, 21.38, left_w, 1.28, name="METHODS_R3VAEBand")
    add_textbox(
        slide,
        x + 0.15,
        21.56,
        left_w - 0.30,
        0.84,
        "embedding  →  anchor  →  residual Q",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_R3VAEScheme",
    )
    add_inline_statement(
        slide,
        x,
        22.80,
        left_w,
        1.34,
        "Зачем",
        "усилить семантическую структуру кодов.",
        name="METHODS_R3VAEWhy",
        label_color=GREEN_DARK,
    )

    add_textbox(
        slide,
        right_x,
        19.32,
        left_w,
        0.56,
        "LogQ",
        size=34,
        bold=True,
        color=GREEN_DARK,
        name="METHODS_LogQHead",
    )
    add_inline_statement(
        slide,
        right_x,
        20.04,
        left_w,
        1.20,
        "Что меняем",
        "поправка logits по частоте SID-префиксов / объектов.",
        name="METHODS_LogQWhat",
    )
    add_band(slide, right_x, 21.38, left_w, 1.28, name="METHODS_LogQBand")
    add_textbox(
        slide,
        right_x + 0.15,
        21.52,
        left_w - 0.30,
        0.96,
        "raw logits  →  frequency correction\n→  corrected logits",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_LogQScheme",
        line_spacing=1.0,
    )
    add_inline_statement(
        slide,
        right_x,
        22.80,
        left_w,
        1.34,
        "Зачем",
        "компенсировать частотный перекос при обучении.",
        name="METHODS_LogQWhy",
        label_color=GREEN_DARK,
    )


def add_generation_table(slide) -> None:
    x, w = 31.19, 13.78
    col_model, col_recall = 7.25, 3.35
    col_delta = w - col_model - col_recall
    header_y, header_h, row_h = 19.10, 0.70, 0.92

    add_textbox(
        slide,
        x,
        18.45,
        w,
        0.48,
        "Главный результат генерации",
        size=34,
        bold=True,
        name="RESULTS_GenerationTableHead",
    )
    add_band(slide, x, header_y, w, header_h, fill=PURPLE_TINT, name="RESULTS_TableHeaderBand")
    add_band(
        slide,
        x,
        header_y + header_h + row_h,
        w,
        row_h,
        fill=PURPLE_TINT,
        name="RESULTS_TableBestBand",
    )

    columns = [
        (x, col_model, "Модель", PP_ALIGN.LEFT),
        (x + col_model, col_recall, "Recall@100", PP_ALIGN.RIGHT),
        (x + col_model + col_recall, col_delta, "Δ", PP_ALIGN.RIGHT),
    ]
    for index, (cell_x, cell_w, label, align) in enumerate(columns):
        add_textbox(
            slide,
            cell_x,
            header_y,
            cell_w,
            header_h,
            label,
            size=30,
            bold=True,
            color=MUTED,
            align=align,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"RESULTS_TableHeader_{index}",
            margins=0.10,
        )

    rows = [
        ("SASRec-Audio-Uniform", "0.12303", "baseline", False),
        ("RQ-VAE + item LogQ", "0.14138", "+14.9%", True),
    ]
    for row_index, (model, recall, delta, is_best) in enumerate(rows):
        y = header_y + header_h + row_index * row_h
        values = [
            (x, col_model, model, PP_ALIGN.LEFT),
            (x + col_model, col_recall, recall, PP_ALIGN.RIGHT),
            (x + col_model + col_recall, col_delta, delta, PP_ALIGN.RIGHT),
        ]
        for col_index, (cell_x, cell_w, value, align) in enumerate(values):
            add_textbox(
                slide,
                cell_x,
                y,
                cell_w,
                row_h,
                value,
                size=36 if is_best and col_index > 0 else 30,
                bold=is_best,
                color=GREEN_DARK if is_best and col_index > 0 else INK,
                align=align,
                valign=MSO_VERTICAL_ANCHOR.MIDDLE,
                name=f"RESULTS_TableCell_{row_index}_{col_index}",
                margins=0.10,
            )
        add_rule(
            slide,
            x,
            y + row_h,
            x + w,
            y + row_h,
            color=PURPLE_RULE,
            width=1,
            name=f"RESULTS_TableRule_{row_index}",
        )

    add_textbox(
        slide,
        x,
        21.78,
        w,
        0.58,
        "Лучший генеративный вариант — RQ-VAE + item LogQ.",
        size=30,
        bold=True,
        color=GREEN_DARK,
        name="RESULTS_GenerationTableNote",
    )


def build_results(slide) -> None:
    x, w = 31.19, 13.78
    add_textbox(
        slide,
        x,
        9.50,
        w,
        1.08,
        "КАКИЕ СВОЙСТВА SID СВЯЗАНЫ\nС DOWNSTREAM-КАЧЕСТВОМ?",
        size=34,
        bold=True,
        name="RESULTS_HeatmapQuestion",
        line_spacing=1.0,
    )

    shapes = {shape.name: shape for shape in slide.shapes}
    heatmap = shapes["RESULTS_Heatmap_Full"]
    heatmap.left = Inches(30.84)
    heatmap.top = Inches(10.72)
    heatmap.width = Inches(14.48)
    heatmap.height = Inches(5.53)

    add_inline_statement(
        slide,
        x,
        16.42,
        w,
        0.72,
        "DCNv2 Pairwise",
        "ICR и code entropy дают r ≈ 0.47.",
        name="RESULTS_Observation_1",
        label_color=GREEN_DARK,
    )
    add_inline_statement(
        slide,
        x,
        17.16,
        w,
        0.72,
        "DCNv2 ROC-AUC",
        "выраженной связи нет, max |r| ≈ 0.20.",
        name="RESULTS_Observation_2",
        label_color=GREEN_DARK,
    )
    add_rule(
        slide,
        x,
        18.06,
        x + w,
        18.06,
        color=PURPLE_RULE,
        width=1.5,
        name="RESULTS_ObservationRule",
    )
    add_generation_table(slide)


def build_conclusions(slide) -> None:
    x_positions = [1.96, 16.28, 30.60]
    col_w = 13.30
    for index in (1, 2):
        divider_x = x_positions[index] - 0.52
        add_rule(
            slide,
            divider_x,
            28.82,
            divider_x,
            31.08,
            color=PURPLE_RULE,
            width=1.5,
            name=f"CONCLUSIONS_Divider_{index}",
        )

    theses = [
        [
            {"runs": [("01", 32, True, GREEN_DARK)]},
            {
                "runs": [
                    (
                        "Внутреннее качество SID само по себе не гарантирует "
                        "downstream-улучшение.",
                        34,
                        True,
                        INK,
                    )
                ],
                "before": 5,
                "line_spacing": 1.0,
            },
        ],
        [
            {"runs": [("02", 32, True, GREEN_DARK)]},
            {"runs": [("Лучший генеративный вариант:", 34, True, INK)], "before": 5},
            {
                "runs": [("Recall@100 0.12303 → 0.14138 (+14.9%).", 34, True, GREEN_DARK)],
                "before": 2,
                "line_spacing": 1.0,
            },
        ],
        [
            {"runs": [("03", 32, True, GREEN_DARK)]},
            {
                "runs": [
                    (
                        "Связь характеристик SID с итоговым качеством зависит от "
                        "downstream-задачи и метрики.",
                        34,
                        True,
                        INK,
                    )
                ],
                "before": 5,
                "line_spacing": 1.0,
            },
        ],
    ]
    for index, paragraphs in enumerate(theses):
        add_rich_textbox(
            slide,
            x_positions[index],
            28.72,
            col_w,
            2.45,
            paragraphs,
            name=f"CONCLUSIONS_Thesis_{index + 1}",
        )


def assert_invariants(prs) -> None:
    assert len(prs.slides) == 1
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    names = {shape.name for shape in slide.shapes}
    required = {
        "INTRO_RQKMeansBand",
        "INTRO_ResearchQuestion",
        "METHODS_UserAwareBand",
        "METHODS_PSIBand",
        "METHODS_R3VAEBand",
        "METHODS_LogQBand",
        "RESULTS_Heatmap_Full",
        "RESULTS_TableBestBand",
        "CONCLUSIONS_Thesis_3",
    }
    assert required <= names, f"Missing shapes: {sorted(required - names)}"
    assert "RESULTS_Recall_vs_ICR_SVG" not in names
    assert "RESULTS_TransAct_vs_IntraCode_SVG" not in names
    assert len(list(slide._element.iter(f"{{{A14_NS}}}m"))) == 0

    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        assert run.font.size.pt >= 30, (
                            f"Text below 30 pt in {shape.name}: {run.text!r} ({run.font.size.pt})"
                        )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    slide = prs.slides[0]
    delete_named_content(slide, ("INTRO_", "METHODS_", "CONCLUSIONS_"))
    delete_named_content(slide, ("RESULTS_",), keep={"RESULTS_Heatmap_Full"})
    delete_hidden_named_content(slide, ("INTRO_", "METHODS_", "CONCLUSIONS_", "RESULTS_"))
    delete_residual_center_content(slide)

    build_introduction(slide)
    build_methods(slide)
    build_results(slide)
    build_conclusions(slide)

    prs.core_properties.comments = (
        "Hallmark · Narrative Workflow · conversation-starter scientific poster · "
        "fixed ML-Академия A0 template · no scientific values changed"
    )
    assert_invariants(prs)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
