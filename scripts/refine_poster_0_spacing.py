from __future__ import annotations

# Hallmark · pre-emit critique: P5 H5 E5 S5 R5 V4
# Constrained single-page redesign · technical/austere · whitespace pass

from hashlib import sha256
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from build_poster_pptx import GREEN_DARK, INK, MUTED


SOURCE = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_0_improved.pptx"
)
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_0_improved_spaced.pptx"
)

EMU_PER_INCH = 914400
FONT_NAME = "Arial"


def shape_map(slide):
    return {shape.name: shape for shape in slide.shapes}


def set_box(shape, x: float, y: float, w: float, h: float) -> None:
    shape.left = Inches(x)
    shape.top = Inches(y)
    shape.width = Inches(w)
    shape.height = Inches(h)


def set_paragraph_defaults(paragraph, *, align=PP_ALIGN.LEFT, line_spacing=1.12) -> None:
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)


def set_run(run, *, size: float, bold: bool, color) -> None:
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = color


def rewrite_plain(
    shape,
    text: str,
    *,
    size: float = 30,
    bold: bool = False,
    color=INK,
    align=PP_ALIGN.LEFT,
    valign=MSO_VERTICAL_ANCHOR.TOP,
    line_spacing: float = 1.12,
    margins: float = 0.04,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(margins)
    tf.margin_right = Inches(margins)
    tf.margin_top = Inches(margins)
    tf.margin_bottom = Inches(margins)
    paragraph = tf.paragraphs[0]
    set_paragraph_defaults(paragraph, align=align, line_spacing=line_spacing)
    run = paragraph.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color)


def rewrite_inline(
    shape,
    label: str,
    body: str,
    *,
    label_color=MUTED,
    size: float = 30,
    line_spacing: float = 1.12,
) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.04)
    tf.margin_right = Inches(0.04)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    paragraph = tf.paragraphs[0]
    set_paragraph_defaults(paragraph, line_spacing=line_spacing)
    label_run = paragraph.add_run()
    label_run.text = f"{label} · "
    set_run(label_run, size=size, bold=True, color=label_color)
    body_run = paragraph.add_run()
    body_run.text = body
    set_run(body_run, size=size, bold=False, color=INK)


def rewrite_paragraphs(shape, paragraphs: list[dict], *, margins: float = 0.04) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.TOP
    tf.margin_left = Inches(margins)
    tf.margin_right = Inches(margins)
    tf.margin_top = Inches(margins)
    tf.margin_bottom = Inches(margins)
    for index, spec in enumerate(paragraphs):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        set_paragraph_defaults(
            paragraph,
            align=spec.get("align", PP_ALIGN.LEFT),
            line_spacing=spec.get("line_spacing", 1.12),
        )
        paragraph.space_before = Pt(spec.get("before", 0))
        run = paragraph.add_run()
        run.text = spec["text"]
        set_run(
            run,
            size=spec.get("size", 30),
            bold=spec.get("bold", False),
            color=spec.get("color", INK),
        )


def remove_shape(slide, name: str) -> None:
    for shape in slide.shapes:
        if shape.name == name:
            shape._element.getparent().remove(shape._element)
            return


def refine_introduction(shapes) -> None:
    x, w = 1.86, 12.79

    set_box(shapes["INTRO_DefinitionHead"], x, 11.98, w, 0.48)
    set_box(shapes["INTRO_Definition"], x, 12.58, w, 0.72)
    rewrite_plain(
        shapes["INTRO_Definition"],
        "SID — дискретные коды объектов для рекомендаций.",
        size=30,
        bold=True,
        line_spacing=1.10,
    )

    set_box(shapes["INTRO_RQKMeansBand"], x, 13.58, w, 2.42)
    set_box(shapes["INTRO_RQKMeansHead"], x + 0.18, 13.68, w - 0.36, 0.44)
    for index in range(5):
        stage = shapes[f"INTRO_RQKMeansStage_{index}"]
        stage.top = Inches(14.51)
    for index in range(4):
        arrow = shapes[f"INTRO_RQKMeansArrow_{index}"]
        arrow.top = Inches(14.60)
    set_box(shapes["INTRO_RQKMeansSequence"], x + 0.20, 15.43, w - 0.40, 0.62)

    set_box(shapes["INTRO_GoalHead"], x, 16.38, w, 0.48)
    set_box(shapes["INTRO_GoalBody"], x, 16.98, w, 1.20)
    rewrite_paragraphs(
        shapes["INTRO_GoalBody"],
        [
            {"text": "• Сравнить способы построения SID.", "size": 30, "bold": True},
            {
                "text": "• Проверить связь SID с downstream-качеством.",
                "size": 30,
                "bold": True,
                "before": 7,
            },
        ],
    )

    set_box(shapes["INTRO_ContextHead"], x, 18.58, w, 0.48)
    set_box(shapes["INTRO_Dataset"], x, 19.16, w, 0.52)

    set_box(shapes["INTRO_TaskDivider"], x + 6.22, 20.02, 0, 4.36)
    set_box(shapes["INTRO_GenerationHead"], x, 20.02, 5.92, 0.50)
    set_box(shapes["INTRO_GenerationBody"], x, 20.64, 5.92, 2.76)
    rewrite_paragraphs(
        shapes["INTRO_GenerationBody"],
        [
            {"text": "TIGER → SID", "size": 30, "bold": True},
            {"text": "SASRec baseline · Recall@100", "size": 30, "before": 8},
        ],
    )

    set_box(shapes["INTRO_RankingHead"], x + 6.52, 20.02, 6.27, 0.50)
    set_box(shapes["INTRO_RankingBody"], x + 6.52, 20.64, 6.27, 2.98)
    rewrite_paragraphs(
        shapes["INTRO_RankingBody"],
        [
            {"text": "DCNv2 / TransAct-like", "size": 30, "bold": True},
            {"text": "SID-признак · ROC-AUC · Pairwise Accuracy", "size": 30, "before": 8},
        ],
    )


def refine_methods(slide, shapes) -> None:
    x, w = 15.78, 13.78
    half_w, gap = 6.58, 0.62
    right_x = x + half_w + gap

    rewrite_inline(
        shapes["METHODS_PLUM"],
        "PLUM-like",
        "меньше кодов + случайный prefix.",
        label_color=GREEN_DARK,
    )

    set_box(shapes["METHODS_RQVAEHead"], x, 11.04, 2.20, 0.50)
    set_box(shapes["METHODS_RQVAENote"], x + 2.20, 11.04, w - 2.20, 0.56)
    picture = shapes["Рисунок 4"]
    set_box(picture, 18.20, 11.58, 8.93, 4.00)

    set_box(shapes["METHODS_UserAwareHead"], x, 15.95, w, 0.56)
    set_box(shapes["METHODS_UserAwareWhat"], x, 16.68, w, 0.60)
    rewrite_inline(
        shapes["METHODS_UserAwareWhat"],
        "Что меняем",
        "учитываем user preference при построении SID.",
    )
    set_box(shapes["METHODS_UserAwareBand"], x, 17.50, w, 1.00)
    for index in range(4):
        set_box(
            shapes[f"METHODS_UserAwareStage_{index}"],
            shapes[f"METHODS_UserAwareStage_{index}"].left / EMU_PER_INCH,
            17.59,
            shapes[f"METHODS_UserAwareStage_{index}"].width / EMU_PER_INCH,
            0.82,
        )
    for index in range(3):
        arrow = shapes[f"METHODS_UserAwareArrow_{index}"]
        arrow.top = Inches(17.72)
    set_box(shapes["METHODS_UserAwareWhy"], x, 18.70, w, 0.56)
    rewrite_inline(
        shapes["METHODS_UserAwareWhy"],
        "Зачем",
        "добавить пользовательскую семантику.",
        label_color=GREEN_DARK,
    )

    remove_shape(slide, "METHODS_UserAwareRule")
    shapes = shape_map(slide)

    set_box(shapes["METHODS_PSIHead"], x, 19.63, w, 0.56)
    set_box(shapes["METHODS_PSIWhat"], x, 20.36, w, 0.60)
    rewrite_inline(
        shapes["METHODS_PSIWhat"],
        "Что меняем",
        "занятый SID → следующая близкая свободная комбинация.",
    )
    set_box(shapes["METHODS_PSIBand"], x, 21.18, w, 1.00)
    set_box(shapes["METHODS_PSIStage_0"], x + 0.28, 21.27, 4.62, 0.82)
    rewrite_plain(
        shapes["METHODS_PSIStage_0"],
        "nearest SID\nзанят  ×",
        size=30,
        bold=True,
        color=MUTED,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        line_spacing=1.0,
    )
    set_box(shapes["METHODS_PSIArrow_0"], x + 5.10, 21.40, 0.34, 0.55)
    set_box(shapes["METHODS_PSIStage_2"], x + 5.70, 21.27, 7.74, 0.82)
    rewrite_plain(
        shapes["METHODS_PSIStage_2"],
        "next closest free SID  ✓",
        size=30,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        line_spacing=1.0,
    )
    remove_shape(slide, "METHODS_PSIStage_1")
    remove_shape(slide, "METHODS_PSIArrow_1")
    remove_shape(slide, "METHODS_PSIRule")
    shapes = shape_map(slide)

    set_box(shapes["METHODS_PSIWhy"], x, 22.38, w, 0.56)
    rewrite_inline(
        shapes["METHODS_PSIWhy"],
        "Зачем",
        "убрать коллизии без нового токена.",
        label_color=GREEN_DARK,
    )

    for prefix, column_x, what_body in (
        ("METHODS_R3VAE", x, "semantic anchor перед residual Q."),
        ("METHODS_LogQ", right_x, "frequency correction logits."),
    ):
        set_box(shapes[f"{prefix}Head"], column_x, 23.24, half_w, 0.50)
        set_box(shapes[f"{prefix}What"], column_x, 23.90, half_w, 0.60)
        rewrite_inline(shapes[f"{prefix}What"], "Что меняем", what_body)
        set_box(shapes[f"{prefix}Band"], column_x, 24.70, half_w, 0.84)
        set_box(shapes[f"{prefix}Why"], column_x, 25.74, half_w, 0.54)

    set_box(shapes["METHODS_R3VAEScheme"], x + 0.12, 24.80, half_w - 0.24, 0.64)
    set_box(shapes["METHODS_LogQScheme"], right_x + 0.12, 24.76, half_w - 0.24, 0.72)
    rewrite_inline(
        shapes["METHODS_R3VAEWhy"],
        "Зачем",
        "структурировать коды.",
        label_color=GREEN_DARK,
    )
    rewrite_inline(
        shapes["METHODS_LogQWhy"],
        "Зачем",
        "снизить frequency bias.",
        label_color=GREEN_DARK,
    )


def refine_results(shapes) -> None:
    x, w = 31.19, 13.78
    set_box(shapes["RESULTS_Observation_1"], x, 16.54, w, 0.64)
    set_box(shapes["RESULTS_Observation_2"], x, 17.36, w, 0.64)
    set_box(shapes["RESULTS_ObservationRule"], x, 18.24, w, 0)

    set_box(shapes["RESULTS_GenerationTableHead"], x, 18.70, w, 0.50)
    header_y, header_h, row_h = 19.42, 0.78, 1.05
    set_box(shapes["RESULTS_TableHeaderBand"], x, header_y, w, header_h)
    set_box(shapes["RESULTS_TableBestBand"], x, header_y + header_h + row_h, w, row_h)

    columns = [(x, 7.25), (x + 7.25, 3.35), (x + 10.60, 3.18)]
    for index, (cell_x, cell_w) in enumerate(columns):
        cell = shapes[f"RESULTS_TableHeader_{index}"]
        set_box(cell, cell_x, header_y, cell_w, header_h)
        tf = cell.text_frame
        tf.margin_left = Inches(0.14)
        tf.margin_right = Inches(0.14)
        tf.margin_top = Inches(0.14)
        tf.margin_bottom = Inches(0.14)

    for row_index in range(2):
        row_y = header_y + header_h + row_index * row_h
        for col_index, (cell_x, cell_w) in enumerate(columns):
            cell = shapes[f"RESULTS_TableCell_{row_index}_{col_index}"]
            set_box(cell, cell_x, row_y, cell_w, row_h)
            tf = cell.text_frame
            tf.margin_left = Inches(0.14)
            tf.margin_right = Inches(0.14)
            tf.margin_top = Inches(0.14)
            tf.margin_bottom = Inches(0.14)
        set_box(shapes[f"RESULTS_TableRule_{row_index}"], x, row_y + row_h, w, 0)

    set_box(shapes["RESULTS_GenerationTableNote"], x, 22.72, w, 0.62)


def refine_conclusions(shapes) -> None:
    x_positions = [2.12, 12.79, 23.46, 34.13]
    col_w = 9.52
    theses = [
        (
            "SID ≠ downstream quality",
            "Внутренние метрики SID не гарантируют downstream-улучшение.",
        ),
        (
            "Ranking: устойчивого улучшения нет",
            "SID-конфигурации не превзошли item-only baseline.",
        ),
        (
            "Generation: +14.9%",
            "RQ-VAE + item LogQ дал лучший Recall@100.",
        ),
        (
            "Проверять нужно downstream",
            "Оценивать SID нужно через реальные downstream-задачи.",
        ),
    ]
    for index, ((head, body), x) in enumerate(zip(theses, x_positions), start=1):
        shape = shapes[f"CONCLUSIONS_Thesis_{index}"]
        set_box(shape, x, 28.84, col_w, 2.18)
        rewrite_paragraphs(
            shape,
            [
                {
                    "text": head,
                    "size": 34,
                    "bold": True,
                    "color": GREEN_DARK,
                    "line_spacing": 1.0,
                },
                {
                    "text": body,
                    "size": 30,
                    "before": 12,
                    "line_spacing": 1.12,
                },
            ],
            margins=0.04,
        )


def assert_invariants(prs, *, rqvae_hash: str, heatmap_hash: str) -> None:
    assert len(prs.slides) == 1
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    shapes = shape_map(slide)
    assert sha256(shapes["Рисунок 4"].image.blob).hexdigest() == rqvae_hash
    assert sha256(shapes["RESULTS_Heatmap_Full"].image.blob).hexdigest() == heatmap_hash
    assert "RESULTS_Recall_vs_ICR_SVG" not in shapes

    required = {
        "INTRO_RQKMeansBand",
        "METHODS_UserAwareBand",
        "METHODS_PSIBand",
        "METHODS_R3VAEBand",
        "METHODS_LogQBand",
        "RESULTS_TableBestBand",
        "CONCLUSIONS_Thesis_4",
    }
    assert required <= shapes.keys()

    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"
        if not shape.name.startswith(("INTRO_", "METHODS_", "RESULTS_", "CONCLUSIONS_")):
            continue
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        assert run.font.size.pt >= 30, (
                            f"Text below 30 pt in {shape.name}: "
                            f"{run.text!r} ({run.font.size.pt})"
                        )


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    slide = prs.slides[0]
    shapes = shape_map(slide)
    rqvae_hash = sha256(shapes["Рисунок 4"].image.blob).hexdigest()
    heatmap_hash = sha256(shapes["RESULTS_Heatmap_Full"].image.blob).hexdigest()

    refine_introduction(shapes)
    refine_methods(slide, shape_map(slide))
    refine_results(shape_map(slide))
    refine_conclusions(shape_map(slide))

    prs.core_properties.comments = (
        "Hallmark · pre-emit P5 H5 E5 S5 R5 V4 · whitespace refinement only · "
        "poster_0_improved.pptx used as the sole base · original RQ-VAE and heatmap "
        "images retained byte-for-byte"
    )
    assert_invariants(prs, rqvae_hash=rqvae_hash, heatmap_hash=heatmap_hash)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
