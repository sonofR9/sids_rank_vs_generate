from __future__ import annotations

# Hallmark · pre-emit critique: P5 H5 E5 S5 R5 V4
# Scientific poster · technical/austere · fixed-template three-column narrative

from pathlib import Path

from pptx import Presentation
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_poster_pptx import (
    GREEN,
    GREEN_DARK,
    INK,
    MUTED,
    PURPLE_RULE,
    add_band,
    add_flow_arrow,
    add_office_math_textbox,
    add_rich_textbox,
    add_rule,
    add_textbox,
    office_math_fraction,
    office_math_run,
    office_math_subscript,
)


SOURCE = Path(
    "/Users/konstantinryadinskiy/Desktop/"
    "poster_semantic_id_final_visual_results.pptx"
)
OUTPUT = Path(
    "/Users/konstantinryadinskiy/Desktop/sids_rank_vs_generate/"
    "deliverables/poster_semantic_id_final_visual_results_polished.pptx"
)

EMU_PER_INCH = 914400
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
A14_NS = "http://schemas.microsoft.com/office/drawing/2010/main"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def delete_visible_shapes(slide, prefixes: tuple[str, ...], *, keep: set[str] | None = None) -> None:
    """Delete only named poster content shapes; the fixed template remains untouched."""
    keep = keep or set()
    for shape in list(slide.shapes):
        if shape.name in keep:
            continue
        if shape.name.startswith(prefixes):
            shape._element.getparent().remove(shape._element)


def delete_alternate_content_shapes(slide, names: set[str]) -> None:
    """Remove Office Math wrappers that python-pptx does not expose as shapes."""
    wrappers = set()
    c_nv_pr_tag = f"{{{P_NS}}}cNvPr"
    alternate_tag = f"{{{MC_NS}}}AlternateContent"
    for node in slide._element.iter(c_nv_pr_tag):
        if node.get("name") not in names:
            continue
        parent = node.getparent()
        while parent is not None and parent.tag != alternate_tag:
            parent = parent.getparent()
        if parent is not None:
            wrappers.add(parent)
    for wrapper in wrappers:
        wrapper.getparent().remove(wrapper)


def build_introduction(slide) -> None:
    x, w = 1.86, 12.79

    add_textbox(
        slide,
        x,
        9.50,
        w,
        0.50,
        "Semantic ID и исследовательский вопрос",
        size=34,
        bold=True,
        name="INTRO_Subhead",
    )
    add_textbox(
        slide,
        x,
        10.14,
        w,
        0.98,
        "Semantic IDs — дискретные коды объектов в рекомендациях.\n"
        "Качество SID не гарантирует downstream-качество.",
        size=30,
        name="INTRO_Definition",
        line_spacing=1.05,
    )
    add_rich_textbox(
        slide,
        x,
        11.32,
        w,
        2.58,
        [
            {
                "runs": [
                    (
                        "Какие свойства Semantic ID действительно отражаются на "
                        "downstream-метриках рекомендательной системы?",
                        42,
                        True,
                        INK,
                    )
                ],
                "line_spacing": 1.0,
            }
        ],
        name="INTRO_ResearchQuestion",
    )
    add_rule(
        slide,
        x,
        14.12,
        x + 7.60,
        14.12,
        color=GREEN,
        width=4,
        name="INTRO_QuestionUnderline",
    )

    add_textbox(slide, x, 14.34, w, 0.50, "Цель", size=34, bold=True, name="INTRO_GoalHead")
    add_rich_textbox(
        slide,
        x,
        14.90,
        w,
        1.32,
        [
            {"runs": [("• Сравнить способы построения SID.", 30, False, INK)]},
            {
                "runs": [
                    (
                        "• Проверить связь SID-метрик с downstream-качеством.",
                        30,
                        False,
                        INK,
                    )
                ],
                "before": 4,
            },
        ],
        name="INTRO_GoalBody",
    )

    add_textbox(
        slide,
        x,
        16.34,
        w,
        0.50,
        "Экспериментальный контекст",
        size=34,
        bold=True,
        name="INTRO_ContextHead",
    )
    rows = [
        ("Dataset", "Yambda-50M, time split"),
        ("Generation", "TIGER vs SASRec"),
        ("Ranking", "DCNv2 / TransAct-like"),
        ("Metrics", "Recall@100 · ROC-AUC · Pairwise Accuracy"),
    ]
    y = 16.92
    for index, (label, value) in enumerate(rows):
        add_textbox(
            slide,
            x,
            y,
            3.35,
            0.62,
            label,
            size=30,
            bold=True,
            color=MUTED,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"INTRO_ContextLabel_{index}",
        )
        add_textbox(
            slide,
            x + 3.55,
            y,
            w - 3.55,
            0.62,
            value,
            size=30,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"INTRO_ContextValue_{index}",
        )
        if index < len(rows) - 1:
            add_rule(
                slide,
                x,
                y + 0.62,
                x + w,
                y + 0.62,
                color=PURPLE_RULE,
                width=1,
                name=f"INTRO_ContextRule_{index}",
            )
        y += 0.68

    add_textbox(
        slide,
        x,
        19.92,
        w,
        0.50,
        "Ключевые SID-метрики",
        size=34,
        bold=True,
        name="INTRO_MetricsHead",
    )
    add_textbox(
        slide,
        x,
        20.52,
        w,
        0.48,
        "Independent Code Rate (ICR)",
        size=32,
        bold=True,
        color=GREEN_DARK,
        name="INTRO_ICR_Head",
    )
    icr_math = (
        office_math_run("ICR", 36, normal=True)
        + office_math_run(" = ", 36, normal=True)
        + office_math_fraction(
            office_math_run("unique used codes", 36, normal=True),
            office_math_subscript("n", "items", 36),
        )
    )
    add_office_math_textbox(
        slide,
        x,
        20.99,
        10.50,
        1.08,
        icr_math,
        name="INTRO_ICR_OfficeMath",
    )
    add_textbox(
        slide,
        x,
        22.10,
        w,
        0.68,
        "Чем ближе к 1, тем меньше объектов разделяют один SID.",
        size=30,
        name="INTRO_ICR_Note",
    )
    add_rich_textbox(
        slide,
        x,
        22.94,
        w,
        1.38,
        [
            {"runs": [("Intra-code Similarity", 32, True, GREEN_DARK)]},
            {
                "runs": [
                    (
                        "Чем выше значение, тем семантически однороднее объекты внутри кода.",
                        30,
                        False,
                        INK,
                    )
                ],
                "before": 4,
                "line_spacing": 1.02,
            },
        ],
        name="INTRO_IntraCode",
    )


def build_methods(slide) -> None:
    x, w = 15.78, 13.78

    add_textbox(
        slide,
        x,
        9.50,
        w,
        0.52,
        "Логика эксперимента",
        size=34,
        bold=True,
        name="METHODS_LogicHead",
    )
    add_band(slide, x, 10.12, w, 1.18, name="METHODS_LogicBand")
    add_textbox(
        slide,
        x + 0.18,
        10.22,
        3.28,
        0.98,
        "Построение\nSemantic ID",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_LogicSID",
    )
    add_flow_arrow(slide, x + 3.50, 10.43, name="METHODS_LogicArrow1")
    add_textbox(
        slide,
        x + 3.95,
        10.22,
        4.18,
        0.98,
        "TIGER / DCNv2 /\nTransAct-like",
        size=30,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_LogicModels",
    )
    add_flow_arrow(slide, x + 8.18, 10.43, name="METHODS_LogicArrow2")
    add_textbox(
        slide,
        x + 8.62,
        10.22,
        4.92,
        0.98,
        "downstream-метрики",
        size=32,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_LogicMetrics",
    )

    add_textbox(
        slide,
        x,
        11.62,
        w,
        0.50,
        "Balanced RQ-KMeans",
        size=34,
        bold=True,
        name="METHODS_BalancedHead",
    )
    capacity_math = (
        office_math_run("w", 38, italic=True)
        + office_math_run(" = ", 38, normal=True)
        + office_math_fraction(
            office_math_run("|V|", 38, italic=True),
            office_math_run("K", 38, italic=True),
        )
    )
    add_office_math_textbox(
        slide,
        x,
        12.18,
        3.40,
        1.08,
        capacity_math,
        name="METHODS_BalancedFormula",
    )
    add_rule(
        slide,
        x + 3.62,
        12.22,
        x + 3.62,
        13.15,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_BalancedDivider",
    )
    add_textbox(
        slide,
        x + 3.92,
        12.20,
        w - 3.92,
        0.98,
        "Capacity задаёт лимит объектов на кластер.",
        size=30,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="METHODS_BalancedNote",
    )

    add_textbox(
        slide,
        x,
        13.56,
        w,
        0.50,
        "RQ-VAE",
        size=36,
        bold=True,
        name="METHODS_RQVAEHead",
    )
    add_band(slide, x, 14.12, w, 1.58, name="METHODS_RQVAEBand")
    rq_labels = ["Encoder", "Residual\ncodebooks 1…L", "Quantized\nrepresentation", "Decoder"]
    rq_widths = [2.18, 4.05, 3.76, 2.18]
    rx = x
    for index, (label, box_width) in enumerate(zip(rq_labels, rq_widths)):
        add_textbox(
            slide,
            rx,
            14.30,
            box_width,
            1.22,
            label,
            size=32,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"METHODS_RQVAEStage_{index + 1}",
        )
        if index < 3:
            add_flow_arrow(
                slide,
                rx + box_width + 0.03,
                14.65,
                w=0.32,
                name=f"METHODS_RQVAEArrow_{index + 1}",
            )
            rx += box_width + 0.43
    add_textbox(
        slide,
        x,
        15.78,
        w,
        0.72,
        "Остаток квантуется по уровням; SID — последовательность кодов.",
        size=30,
        color=MUTED,
        name="METHODS_RQVAENote",
    )

    add_textbox(
        slide,
        x,
        16.68,
        w,
        0.50,
        "User-aware",
        size=36,
        bold=True,
        name="METHODS_UserAwareHead",
    )
    add_band(slide, x, 17.24, w, 1.64, name="METHODS_UserAwareBand")
    ua_labels = ["User + item\nembedding", "User-aware\nnetwork", "User-aware SID", "Canonical\nSID target"]
    ua_widths = [2.92, 2.94, 3.10, 2.92]
    ux = x
    for index, (label, box_width) in enumerate(zip(ua_labels, ua_widths)):
        add_textbox(
            slide,
            ux,
            17.43,
            box_width,
            1.26,
            label,
            size=32,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_VERTICAL_ANCHOR.MIDDLE,
            name=f"METHODS_UserAwareStage_{index + 1}",
        )
        if index < 3:
            add_flow_arrow(
                slide,
                ux + box_width + 0.06,
                17.80,
                w=0.36,
                name=f"METHODS_UserAwareArrow_{index + 1}",
            )
            ux += box_width + 0.50

    add_rule(
        slide,
        x,
        19.30,
        x + w,
        19.30,
        color=PURPLE_RULE,
        width=1.5,
        name="METHODS_SummaryRule",
    )
    summaries = [
        ("PLUM-like", "Кодов меньше; префикс случаен."),
        ("R3-VAE", "Якорь + SC/PD-регуляризация."),
        ("PSI", "Коллизии без нового токена."),
        ("LogQ", "Частотная коррекция SID."),
    ]
    positions = [(x, 19.70), (x + 7.08, 19.70), (x, 22.00), (x + 7.08, 22.00)]
    for index, ((head, body), (box_x, box_y)) in enumerate(zip(summaries, positions)):
        add_rich_textbox(
            slide,
            box_x,
            box_y,
            6.70,
            1.58,
            [
                {"runs": [(head, 32, True, GREEN_DARK)]},
                {
                    "runs": [(body, 30, False, INK)],
                    "before": 5,
                    "line_spacing": 1.02,
                },
            ],
            name=f"METHODS_Summary_{index + 1}",
        )


def build_generation_highlight(slide) -> None:
    x, w = 31.19, 13.78
    add_rule(
        slide,
        x,
        9.47,
        x + w,
        9.47,
        color=GREEN,
        width=3,
        name="RESULTS_GenerationRule",
    )
    add_textbox(
        slide,
        x,
        9.54,
        w,
        0.40,
        "GENERATION УЛУЧШИЛСЯ · Recall@100",
        size=32,
        bold=True,
        color=GREEN_DARK,
        name="RESULTS_GenerationHead",
        line_spacing=1.0,
    )
    add_textbox(
        slide,
        x,
        9.94,
        8.15,
        0.65,
        "0.12303 → 0.14138",
        size=50,
        bold=True,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="RESULTS_GenerationValues",
        line_spacing=1.0,
    )
    add_textbox(
        slide,
        x + 8.23,
        9.90,
        5.47,
        0.72,
        "+14.9%",
        size=60,
        bold=True,
        color=GREEN_DARK,
        align=PP_ALIGN.RIGHT,
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
        name="RESULTS_GenerationDelta",
        line_spacing=1.0,
    )
    add_rich_textbox(
        slide,
        x,
        10.60,
        w,
        0.43,
        [
            {
                "runs": [
                    ("Baseline · ", 30, True, MUTED),
                    ("SASRec-Audio-Uniform", 30, False, INK),
                ],
                "line_spacing": 1.0,
            }
        ],
        name="RESULTS_GenerationBaseline",
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )
    add_rich_textbox(
        slide,
        x,
        11.02,
        w,
        0.43,
        [
            {
                "runs": [
                    ("Best SID · ", 30, True, MUTED),
                    ("RQ-VAE + item LogQ, L=2, C=256", 30, False, INK),
                ],
                "line_spacing": 1.0,
            }
        ],
        name="RESULTS_GenerationBest",
        valign=MSO_VERTICAL_ANCHOR.MIDDLE,
    )


def enlarge_results_figures(slide) -> None:
    shapes = {shape.name: shape for shape in slide.shapes}
    heatmap = shapes["RESULTS_Heatmap_Full"]
    heatmap.left = Inches(30.98)
    heatmap.top = Inches(11.56)
    heatmap.width = Inches(14.20)
    heatmap.height = Inches(5.42)

    recall = shapes["RESULTS_Recall_vs_ICR_SVG"]
    recall.left = Inches(31.19)
    recall.top = Inches(17.35)
    recall.width = Inches(13.78)
    recall.height = Inches(5.91)


def build_conclusions(slide) -> None:
    theses = [
        (
            "SID ≠ downstream quality",
            "Хорошие внутренние метрики SID не гарантируют улучшение итоговой модели.",
        ),
        (
            "Ranking: улучшения нет",
            "SID-конфигурации не превзошли item-only baseline в ranking-задачах.",
        ),
        (
            "Generation: +14.9%",
            "RQ-VAE + item LogQ улучшил Recall@100 относительно baseline.",
        ),
        (
            "Проверять нужно downstream",
            "SID нужно оценивать не только внутренними метриками, но и downstream-экспериментами.",
        ),
    ]
    x0, col_w, gap = 1.96, 9.85, 0.82
    for index, (head, body) in enumerate(theses):
        x = x0 + index * (col_w + gap)
        if index:
            add_rule(
                slide,
                x - 0.40,
                28.90,
                x - 0.40,
                30.98,
                color=PURPLE_RULE,
                width=1.5,
                name=f"CONCLUSIONS_Divider_{index}",
            )
        add_rich_textbox(
            slide,
            x,
            28.82,
            col_w,
            2.18,
            [
                {"runs": [(head, 36, True, GREEN_DARK)], "line_spacing": 1.0},
                {
                    "runs": [(body, 30, False, INK)],
                    "before": 7,
                    "line_spacing": 1.02,
                },
            ],
            name=f"CONCLUSIONS_Thesis_{index + 1}",
        )


def assert_invariants(prs) -> None:
    assert len(prs.slides) == 1, "Poster must contain exactly one slide"
    assert abs(prs.slide_width / EMU_PER_INCH - 47.0416667) < 0.001
    assert abs(prs.slide_height / EMU_PER_INCH - 33.3333333) < 0.001

    slide = prs.slides[0]
    names = {shape.name for shape in slide.shapes}
    required = {
        "RESULTS_Heatmap_Full",
        "RESULTS_Recall_vs_ICR_SVG",
        "INTRO_ResearchQuestion",
        "METHODS_BalancedFormula",
    }
    assert required <= names, f"Missing required content shapes: {sorted(required - names)}"
    assert "RESULTS_TransAct_vs_IntraCode_SVG" not in names

    for shape in slide.shapes:
        assert shape.left >= 0 and shape.top >= 0, f"Off-slide shape: {shape.name}"
        assert shape.left + shape.width <= prs.slide_width + 2, f"Right overflow: {shape.name}"
        assert shape.top + shape.height <= prs.slide_height + 2, f"Bottom overflow: {shape.name}"
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.text.strip() and run.font.size is not None:
                        assert run.font.size.pt >= 30, (
                            f"Text below 30 pt in {shape.name}: {run.text!r} "
                            f"({run.font.size.pt})"
                        )

    math_objects = list(slide._element.iter(f"{{{A14_NS}}}m"))
    fractions = list(slide._element.iter(f"{{{M_NS}}}f"))
    subscripts = list(slide._element.iter(f"{{{M_NS}}}sSub"))
    assert len(math_objects) == 2, "Poster must contain ICR and Balanced RQ-KMeans equations"
    assert len(fractions) == 2, "Both equations must use native Office Math fractions"
    assert len(subscripts) == 1, "ICR denominator must retain the native items subscript"
    for math_object in math_objects:
        sizes = [
            int(node.get("sz"))
            for node in math_object.iter(f"{{{A_NS}}}rPr")
            if node.get("sz") is not None
        ]
        assert sizes and min(sizes) >= 3000, "Office Math text must be at least 30 pt"


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    prs = Presentation(SOURCE)
    slide = prs.slides[0]

    delete_visible_shapes(slide, ("INTRO_", "METHODS_", "CONCLUSIONS_"))
    delete_visible_shapes(
        slide,
        ("RESULTS_",),
        keep={"RESULTS_Heatmap_Full", "RESULTS_Recall_vs_ICR_SVG"},
    )
    delete_alternate_content_shapes(
        slide,
        {"INTRO_ICR_OfficeMath", "METHODS_BalancedConcept", "METHODS_LogQ_OfficeMath"},
    )

    build_introduction(slide)
    build_methods(slide)
    build_generation_highlight(slide)
    enlarge_results_figures(slide)
    build_conclusions(slide)

    prs.core_properties.comments = (
        "Hallmark · constrained fixed-template audit · technical/austere · "
        "content density reduced; scientific values preserved; no macrostructure redesign"
    )
    assert_invariants(prs)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
